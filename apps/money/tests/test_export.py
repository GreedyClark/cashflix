import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.money.models import Transaction


@pytest.fixture
def yura_client(client, yura):
    client.force_login(yura.user)
    return client


def add(home, account, category, amount, date, type="expense", title="x", **extra):
    return Transaction.objects.create(
        household=home, type=type, account=account, category=category, amount=Decimal(amount), date=date,
        title=title, **extra,
    )


def export(client, **params):
    response = client.get(reverse("money:transaction_export"), params)
    assert response.status_code == 200
    return response, load_workbook(BytesIO(response.content))


def test_export_has_three_sheets_and_signed_amounts(yura_client, home, mono, food, salary):
    add(home, mono, food, "250.50", datetime.date(2026, 9, 3), title="Сільпо")
    add(home, mono, salary, "30000", datetime.date(2026, 9, 1), type="income", title="Зарплата")

    response, book = export(yura_client, month="2026-09")

    assert response["Content-Disposition"] == 'attachment; filename="cashflix-2026-09.xlsx"'
    assert book.sheetnames == ["Операції", "Категорії", "По місяцях"]
    rows = list(book["Операції"].iter_rows(min_row=2, max_row=3, values_only=True))
    assert rows[0][1:3] == ("Дохід", "Зарплата")
    assert rows[1][2] == "Сільпо"
    assert rows[1][7] == -250.5
    assert book["Операції"]["J5"].value == "=SUM(J2:J3)"


def test_export_respects_filters(yura_client, home, mono, food, home_subs):
    add(home, mono, food, "100", datetime.date(2026, 9, 3), title="Сільпо")
    add(home, mono, home_subs, "199", datetime.date(2026, 9, 4), title="Spotify")
    add(home, mono, food, "100", datetime.date(2026, 8, 3), title="Серпень")

    _, book = export(yura_client, month="2026-09", category=food.pk)

    titles = [row[2] for row in book["Операції"].iter_rows(min_row=2, values_only=True) if row[2]]
    assert titles == ["Сільпо"]


def test_categories_sheet_uses_formulas(yura_client, home, mono, food):
    add(home, mono, food, "100", datetime.date(2026, 9, 3))

    _, book = export(yura_client, month="2026-09")

    sheet = book["Категорії"]
    assert sheet["A2"].value == food.name
    assert sheet["B2"].value.startswith("=-SUMIFS(")


def test_transfer_has_no_base_amount(yura_client, home, mono, cash):
    add(home, mono, None, "500", datetime.date(2026, 9, 5), type="transfer", to_account=cash)

    _, book = export(yura_client, month="2026-09")

    row = next(book["Операції"].iter_rows(min_row=2, max_row=2, values_only=True))
    assert row[5] == cash.name
    assert row[9] is None


def test_export_only_own_household(client, user, user_household, home, mono, food):
    add(home, mono, food, "100", datetime.date(2026, 9, 3), title="Чуже")
    client.force_login(user)

    _, book = export(client, month="2026-09")

    assert book["Операції"].max_row == 1
