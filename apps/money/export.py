from decimal import Decimal
from io import BytesIO

from django.db.models import Sum
from django.db.models.functions import TruncMonth
from django.utils import dateformat
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.recurring.dates import add_months

from .models import Transaction

HEADER_FILL = PatternFill("solid", fgColor="EEF0FB")
HEADER_FONT = Font(bold=True, color="1F2330")
MONEY_FORMAT = '#,##0.00;[Red]-#,##0.00'
DATE_FORMAT = "DD.MM.YYYY"
OPERATIONS = "Операції"


FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def plain(text):
    text = text or ""
    return f"'{text}" if text.startswith(FORMULA_START) else text


def type_name(value):
    return str(Transaction.Type(value).label).capitalize()


def signed(tx):
    if tx.type == Transaction.Type.EXPENSE:
        return -tx.amount, -tx.amount_base
    return tx.amount, tx.amount_base


def style_header(sheet, widths):
    for index, width in enumerate(widths, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 22


def operations_sheet(book, transactions, base_currency):
    sheet = book.active
    sheet.title = OPERATIONS
    sheet.append([
        _("Дата"), _("Тип"), _("Назва"), _("Категорія"), _("Рахунок"), _("На рахунок"), _("Хто"),
        _("Сума"), _("Валюта"), _("Сума, %(currency)s") % {"currency": base_currency}, _("Нотатка"),
    ])
    for tx in transactions:
        amount, amount_base = signed(tx)
        sheet.append([
            tx.date,
            type_name(tx.type),
            plain(tx.title),
            plain(tx.category.name) if tx.category else "",
            plain(tx.account.name),
            plain(tx.to_account.name) if tx.to_account else "",
            plain(tx.member.user.name) if tx.member else _("Спільне"),
            amount,
            tx.currency,
            amount_base if tx.type != Transaction.Type.TRANSFER else None,
            plain(tx.note),
        ])
    last = sheet.max_row
    for row in sheet.iter_rows(min_row=2, max_row=last):
        row[0].number_format = DATE_FORMAT
        row[7].number_format = MONEY_FORMAT
        row[9].number_format = MONEY_FORMAT
    if last > 1:
        sheet.append([])
        total_row = last + 2
        sheet.cell(row=total_row, column=9, value=_("Разом")).font = Font(bold=True)
        total = sheet.cell(row=total_row, column=10, value=f"=SUM(J2:J{last})")
        total.font = Font(bold=True)
        total.number_format = MONEY_FORMAT
        sheet.auto_filter.ref = f"A1:K{last}"
    style_header(sheet, [12, 11, 28, 16, 16, 16, 12, 13, 9, 14, 30])
    return last


def categories_sheet(book, last_row):
    sheet = book.create_sheet(_("Категорії"))
    sheet.append([_("Категорія"), _("Витрати"), _("Доходи"), _("Частка витрат")])
    names = sorted({cell.value for cell in book[OPERATIONS]["D"][1:last_row] if cell.value})
    expense, income = type_name(Transaction.Type.EXPENSE), type_name(Transaction.Type.INCOME)
    source = f"'{OPERATIONS}'!$J$2:$J${last_row}"
    kinds = f"'{OPERATIONS}'!$B$2:$B${last_row}"
    categories = f"'{OPERATIONS}'!$D$2:$D${last_row}"
    first = 2
    end = first + len(names) - 1
    for row, name in enumerate(names, start=first):
        sheet.append([
            name,
            f'=-SUMIFS({source},{kinds},"{expense}",{categories},A{row})',
            f'=SUMIFS({source},{kinds},"{income}",{categories},A{row})',
            f"=IF(SUM($B${first}:$B${end})=0,0,B{row}/SUM($B${first}:$B${end}))",
        ])
        sheet[f"B{row}"].number_format = MONEY_FORMAT
        sheet[f"C{row}"].number_format = MONEY_FORMAT
        sheet[f"D{row}"].number_format = "0%"
    if names:
        total_row = end + 1
        sheet.cell(row=total_row, column=1, value=_("Разом")).font = Font(bold=True)
        for column in ("B", "C"):
            cell = sheet[f"{column}{total_row}"]
            cell.value = f"=SUM({column}{first}:{column}{end})"
            cell.font = Font(bold=True)
            cell.number_format = MONEY_FORMAT
    style_header(sheet, [22, 14, 14, 14])


def months_sheet(book, household, today, count=12):
    sheet = book.create_sheet(_("По місяцях"))
    sheet.append([_("Місяць"), _("Дохід"), _("Витрати"), _("Заощаджено"), _("Норма заощаджень")])
    start = add_months(today.replace(day=1), -(count - 1))
    rows = (
        household.transactions.filter(date__gte=start, date__lte=today)
        .exclude(type=Transaction.Type.TRANSFER)
        .annotate(period=TruncMonth("date"))
        .values("period", "type")
        .annotate(total=Sum("amount_base"))
    )
    totals = {(row["period"], row["type"]): row["total"] for row in rows}
    for offset in range(count):
        month = add_months(start, offset)
        row = offset + 2
        sheet.append([
            dateformat.format(month, "F Y"),
            totals.get((month, Transaction.Type.INCOME), Decimal("0")),
            totals.get((month, Transaction.Type.EXPENSE), Decimal("0")),
            f"=B{row}-C{row}",
            f"=IF(B{row}=0,0,D{row}/B{row})",
        ])
        for column in ("B", "C", "D"):
            sheet[f"{column}{row}"].number_format = MONEY_FORMAT
        sheet[f"E{row}"].number_format = "0%"
    style_header(sheet, [18, 14, 14, 14, 18])


def build_workbook(household, transactions, today):
    book = Workbook()
    last_row = operations_sheet(book, transactions, household.base_currency)
    categories_sheet(book, last_row)
    months_sheet(book, household, today)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def file_name(month):
    return f"cashflix-{month:%Y-%m}.xlsx"

